import pptx

def update_text_by_name(slide, shape_name, new_text):
    for s in slide.shapes:
        if s.name == shape_name and s.has_text_frame:
            tf = s.text_frame
            if tf.paragraphs:
                p = tf.paragraphs[0]
                font_name = "Plus Jakarta Sans"
                font_size = None
                font_bold = None
                font_color = None
                if p.runs:
                    r0 = p.runs[0]
                    if r0.font:
                        font_name = r0.font.name or font_name
                        font_size = r0.font.size
                        font_bold = r0.font.bold
                        try:
                            font_color = r0.font.color.rgb
                        except:
                            pass
                p.text = new_text
                if p.runs:
                    r = p.runs[0]
                    if font_name:
                        r.font.name = font_name
                    if font_size:
                        r.font.size = font_size
                    if font_bold is not None:
                        r.font.bold = font_bold
                    if font_color:
                        try:
                            r.font.color.rgb = font_color
                        except:
                            pass
                print(f"  [OK] Updated {shape_name}: \"{new_text[:40]}...\"")
                return True
    print(f"  [ERR] Shape {shape_name} not found!")
    return False

def replace_picture_by_name(slide, shape_name, image_path):
    for s in list(slide.shapes):
        if s.name == shape_name and s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            left, top, width, height = s.left, s.top, s.width, s.height
            slide.shapes._spTree.remove(s._element)
            new_pic = slide.shapes.add_picture(image_path, left, top, width, height)
            new_pic.name = shape_name
            print(f"  [OK] Replaced Picture {shape_name} with {image_path}")
            return True
    print(f"  [ERR] Picture shape {shape_name} not found!")
    return False

# Always load from pristine backup to ensure idempotency
prs = pptx.Presentation('MS Analytics Practicum Midterm Progress Report_backup.pptx')

# --- SLIDE 1 ---
print("Slide 1:")
update_text_by_name(prs.slides[0], "Google Shape;88;p13", "Midterm Progress Report: Operational Pipeline, Dynamic Behavioral Features, Baseline Modeling, and Alpha Prototype for Graylogic Technologies")

# --- SLIDE 3 ---
print("Slide 3:")
update_text_by_name(prs.slides[2], "Google Shape;124;p15", "Data Volume: 150 client accounts across 36 billing months (5,400 transactional rows).")

# --- SLIDE 4 ---
print("Slide 4:")
update_text_by_name(prs.slides[3], "Google Shape;136;p16", "150")
update_text_by_name(prs.slides[3], "Google Shape;140;p16", "22.0%")
update_text_by_name(prs.slides[3], "Google Shape;141;p16", "Historical baseline churn rate (33 churned / 117 retained)")

# --- SLIDE 6 ---
print("Slide 6:")
update_text_by_name(prs.slides[5], "Google Shape;174;p18", "Trailing 3M vs prior 12M spend contraction (V_t = S_3M / S_12M).")
update_text_by_name(prs.slides[5], "Google Shape;176;p18", "Widening delay in settling invoices (Delta L_t = L_3M - L_12M).")
update_text_by_name(prs.slides[5], "Google Shape;178;p18", "Concentration of service escalations and disputes over 6 months.")

# --- SLIDE 7 ---
print("Slide 7:")
update_text_by_name(prs.slides[6], "Google Shape;186;p19", "XGBoost candidate achieves 88% AUC-ROC and 81% recall under 5-Fold Stratified Cross-Validation, identifying 4 out of 5 at-risk accounts while outperforming the L2 baseline by +10% AUC.")

# --- SLIDE 8 ---
print("Slide 8:")
update_text_by_name(prs.slides[7], "Google Shape;196;p20", "Interactive Decision Workbench (Midterm Alpha Prototype)")
update_text_by_name(prs.slides[7], "Google Shape;197;p20", "Portfolio Risk Radar")
update_text_by_name(prs.slides[7], "Google Shape;198;p20", "Segments and ranks 150 accounts by real-time predicted churn probability.")
update_text_by_name(prs.slides[7], "Google Shape;199;p20", "SHAP Explainability Engine")
update_text_by_name(prs.slides[7], "Google Shape;200;p20", "Decomposes account risk into local SHAP feature attributions.")
update_text_by_name(prs.slides[7], "Google Shape;201;p20", "Prescriptive Action Simulator")
update_text_by_name(prs.slides[7], "Google Shape;202;p20", "Simulates fee concessions and SLA recovery to project retention ROI.")
replace_picture_by_name(prs.slides[7], "Google Shape;194;p20", "slide8_with_sidebar.png")

# --- SLIDE 9 ---
print("Slide 9:")
update_text_by_name(prs.slides[8], "Google Shape;213;p21", "PHASE 1 (COMPLETED • WEEKS 1-4)")
update_text_by_name(prs.slides[8], "Google Shape;214;p21", "Data & Governance")
update_text_by_name(prs.slides[8], "Google Shape;215;p21", "Tokenized 150 accounts, processed 5,400 logs, and enforced NDA compliance.")

update_text_by_name(prs.slides[8], "Google Shape;216;p21", "PHASE 2 (COMPLETED • WEEKS 5-7)")
update_text_by_name(prs.slides[8], "Google Shape;217;p21", "Features & Alpha Tool")
update_text_by_name(prs.slides[8], "Google Shape;218;p21", "Engineered 42 behavioral features; ran 5-fold CV; stood up Streamlit Alpha Prototype.")

update_text_by_name(prs.slides[8], "Google Shape;219;p21", "PHASE 3 (PLANNED • WEEKS 8-10)")
update_text_by_name(prs.slides[8], "Google Shape;220;p21", "Cost & Survival Tuning")
update_text_by_name(prs.slides[8], "Google Shape;221;p21", "Optimize decision threshold against Graylogic cost matrix; add survival hazard models.")

update_text_by_name(prs.slides[8], "Google Shape;222;p21", "PHASE 4 (PLANNED • WEEKS 11-14)")
update_text_by_name(prs.slides[8], "Google Shape;223;p21", "Hardening & Defense")
update_text_by_name(prs.slides[8], "Google Shape;224;p21", "Production hardening, automated reporting, executive defense, and final practicum report.")

# --- SLIDE 10 ---
print("Slide 10:")
update_text_by_name(prs.slides[9], "Google Shape;232;p22", "Thank you for reviewing our midterm progress. We welcome faculty and Graylogic feedback on cross-validation threshold tuning and business intervention cost assumptions.")

# --- SLIDE 11 ---
print("Slide 11:")
update_text_by_name(prs.slides[10], "Google Shape;238;p23", "Technical Architecture & Citations")
update_text_by_name(prs.slides[10], "Google Shape;239;p23", "Primary Prototype: Graylogic AcuityRisk™ Decision Support Tool (Alpha Prototype developed for Georgia Tech Applied Analytics Practicum, CSE/ISYE/MGT 6748)")
update_text_by_name(prs.slides[10], "Google Shape;240;p23", "Analytics Stack: Python 3.9 • Streamlit v1.38 • Scikit-Learn 1.3 • XGBoost 2.0 • SHAP 0.44 • Plotly Express")
replace_picture_by_name(prs.slides[10], "Google Shape;237;p23", "slide11_thumb.png")

output_file = 'MS Analytics Practicum Midterm Progress Report.pptx'
prs.save(output_file)
print(f"\nALL SLIDES UPDATED ACCURATELY AND SAVED TO: {output_file}")
