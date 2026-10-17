import pptx

prs = pptx.Presentation('MS Analytics Practicum Midterm Progress Report_backup.pptx')
slide = prs.slides[7] # Slide 8
pic_found = False
for s in list(slide.shapes):
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        left, top, width, height = s.left, s.top, s.width, s.height
        print(f"Replacing pic: {s.name} at {left}, {top}, {width}, {height}")
        slide.shapes._spTree.remove(s._element)
        slide.shapes.add_picture('slide8_with_sidebar.png', left, top, width, height)
        pic_found = True
        break

print("Pic replaced:", pic_found)
prs.save('test_output.pptx')

# Test re-reading
prs2 = pptx.Presentation('test_output.pptx')
print("Successfully re-opened test_output.pptx with slides:", len(prs2.slides))
