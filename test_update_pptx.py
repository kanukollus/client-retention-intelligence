import pptx
from pptx.util import Inches

def update_text(shape, new_text):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        return
    p = tf.paragraphs[0]
    font_name = None
    font_size = None
    font_bold = None
    font_color = None
    if p.runs:
        r0 = p.runs[0]
        if r0.font:
            font_name = r0.font.name
            font_size = r0.font.size
            font_bold = r0.font.bold
            try:
                font_color = r0.font.color.rgb
            except:
                font_color = None
    
    p.text = new_text
    if p.runs and font_name:
        p.runs[0].font.name = font_name
    if p.runs and font_size:
        p.runs[0].font.size = font_size
    if p.runs and font_bold is not None:
        p.runs[0].font.bold = font_bold
    if p.runs and font_color:
        try:
            p.runs[0].font.color.rgb = font_color
        except:
            pass

prs = pptx.Presentation('MS Analytics Practicum Midterm Progress Report_backup.pptx')
print("Successfully loaded backup presentation.")
