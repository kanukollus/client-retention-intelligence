import pptx
from pptx.util import Inches, Pt

def update_shape_text(shape, new_text):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        return
    p = tf.paragraphs[0]
    
    # Preserve formatting from existing runs if present
    font_name = "Plus Jakarta Sans"
    font_size = None
    font_bold = None
    font_color = None
    
    for r in p.runs:
        if r.font:
            if r.font.name:
                font_name = r.font.name
            if r.font.size:
                font_size = r.font.size
            if r.font.bold is not None:
                font_bold = r.font.bold
            try:
                if r.font.color and r.font.color.rgb:
                    font_color = r.font.color.rgb
            except:
                pass
            break
            
    p.text = new_text
    if p.runs:
        r = p.runs[0]
        if font_name:
            r.font.name = font_name
        if font_size:
            r.font.size = font_size
        if font_bold is not None:
            r.font.bold = font_bold
        if font_color:
            try:
                r.font.color.rgb = font_color
            except:
                pass

print("Loading presentation...")
prs = pptx.Presentation('MS Analytics Practicum Midterm Progress Report.pptx')

# --- SLIDE 1 ---
print("Updating Slide 1...")
s1 = prs.slides[0]
update_shape_text(s1.shapes[4], "Midterm Progress Report: Operational Pipeline, Dynamic Behavioral Features, Baseline Modeling, and Alpha Prototype for Graylogic Technologies")

# --- SLIDE 3 ---
print("Updating Slide 3...")
s3 = prs.slides[2]
update_shape_text(s3.shapes[8], "Data Volume: 150 client accounts across 36 billing months (5,400 transactional rows).")

# --- SLIDE 4 ---
print("Updating Slide 4...")
s4 = prs.slides[3]
update_shape_text(s4.shapes[4], "150")
update_shape_text(s4.shapes[8], "22.0%")
update_shape_text(s4.shapes[9], "Historical baseline churn rate (33 churned / 117 retained)")

# --- SLIDE 6 ---
print("Updating Slide 6...")
s6 = prs.slides[5]
update_shape_text(s6.shapes[6], "Trailing 3M vs prior 12M spend contraction (V_t = S_3M / S_12M).")
update_shape_text(s6.shapes[8], "Widening delay in settling invoices (Delta L_t = L_3M - L_12M).")
update_shape_text(s6.shapes[10], "Concentration of service escalations and disputes over 6 months.")

# --- SLIDE 7 ---
print("Updating Slide 7...")
s7 = prs.slides[6]
update_shape_text(s7.shapes[3], "XGBoost candidate achieves 88% AUC-ROC and 81% recall under 5-Fold Stratified Cross-Validation, identifying 4 out of 5 at-risk accounts while outperforming the L2 baseline by +10% AUC.")

# --- SLIDE 8 ---
print("Updating Slide 8...")
s8 = prs.slides[7]
# Replace picture
pic_shape = None
for s in list(s8.shapes):
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        pic_shape = s
        break

if pic_shape:
    left, top, width, height = pic_shape.left, pic_shape.top, pic_shape.width, pic_shape.height
    s8.shapes._spTree.remove(pic_shape._element)
    s8.shapes.add_picture('slide8_with_sidebar.png', left, top, width, height)
    print("Replaced Slide 8 picture with slide8_with_sidebar.png")

update_shape_text(s8.shapes[5], "Interactive Decision Workbench (Midterm Alpha Prototype)")
update_shape_text(s8.shapes[6], "Portfolio Risk Radar")
update_shape_text(s8.shapes[7], "Segments and ranks 150 accounts by real-time predicted churn probability.")
update_shape_text(s8.shapes[8], "SHAP Explainability Engine")
update_shape_text(s8.shapes[9], "Decomposes account risk into local SHAP feature attributions.")
update_shape_text(s8.shapes[10], "Prescriptive Action Simulator")
update_shape_text(s8.shapes[11], "Simulates fee concessions and SLA recovery to project retention ROI.")

# --- SLIDE 9 ---
print("Updating Slide 9...")
s9 = prs.slides[8]
update_shape_text(s9.shapes[6], "PHASE 1 (COMPLETED • WEEKS 1-4)")
update_shape_text(s9.shapes[7], "Data & Governance")
update_shape_text(s9.shapes[8], "Tokenized 150 accounts, processed 5,400 logs, and enforced NDA compliance.")

update_shape_text(s9.shapes[9], "PHASE 2 (COMPLETED • WEEKS 5-7)")
update_shape_text(s9.shapes[10], "Features & Alpha Tool")
update_shape_text(s9.shapes[11], "Engineered 42 behavioral features; ran 5-fold CV; stood up Streamlit Alpha Prototype.")

update_shape_text(s9.shapes[12], "PHASE 3 (PLANNED • WEEKS 8-10)")
update_shape_text(s9.shapes[13], "Cost & Survival Tuning")
update_shape_text(s9.shapes[14], "Optimize decision threshold against Graylogic cost matrix; add survival hazard models.")

update_shape_text(s9.shapes[15], "PHASE 4 (PLANNED • WEEKS 11-14)")
update_shape_text(s9.shapes[16], "Hardening & Defense")
update_shape_text(s9.shapes[17], "Production hardening, automated reporting, executive defense, and final practicum report.")

# --- SLIDE 10 ---
print("Updating Slide 10...")
s10 = prs.slides[9]
update_shape_text(s10.shapes[3], "Thank you for reviewing our midterm progress. We welcome faculty and Graylogic feedback on cross-validation threshold tuning and business intervention cost assumptions.")

# --- SLIDE 11 ---
print("Updating Slide 11...")
s11 = prs.slides[10]
# Replace thumbnail picture
pic11 = None
for s in list(s11.shapes):
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        pic11 = s
        break

if pic11:
    left, top, width, height = pic11.left, pic11.top, pic11.width, pic11.height
    s11.shapes._spTree.remove(pic11._element)
    s11.shapes.add_picture('slide11_thumb.png', left, top, width, height)
    print("Replaced Slide 11 picture with slide11_thumb.png")

update_shape_text(s11.shapes[1], "Technical Architecture & Citations")
update_shape_text(s11.shapes[2], "Primary Prototype: Graylogic AcuityRisk™ Decision Support Tool (Alpha Prototype developed for Georgia Tech Applied Analytics Practicum, CSE/ISYE/MGT 6748)")
update_shape_text(s11.shapes[3], "Analytics Stack: Python 3.9 • Streamlit v1.38 • Scikit-Learn 1.3 • XGBoost 2.0 • SHAP 0.44 • Plotly Express")

# Save updated presentation
output_path = 'MS Analytics Practicum Midterm Progress Report.pptx'
prs.save(output_path)
print(f"Successfully saved updated presentation to {output_path}!")
